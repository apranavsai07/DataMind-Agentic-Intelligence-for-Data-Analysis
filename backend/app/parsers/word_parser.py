from docx import Document as DocxDocument

from app.models.document import Document
from app.parsers.base_parser import BaseParser


class WordParser(BaseParser):

    @staticmethod
    def parse(file_path):

        doc = DocxDocument(file_path)

        text = "\n".join(
            paragraph.text
            for paragraph in doc.paragraphs
            if paragraph.text.strip()
        )

        return Document(
            filename=file_path.name,
            file_type="docx",
            content_type="text",
            content=text,
            metadata={
                "paragraphs": len(doc.paragraphs),
                "characters": len(text),
                "words": len(text.split())
            }
        )